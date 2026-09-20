"""학회 양식 docx 조립.

본문의 **단일 출처는 HTML 원고**다. 이 스크립트는 그것을 읽어 템플릿 위에 옮겨 담는다.
예전에는 같은 본문을 여기에 한 벌 더 들고 있었고, 그 결과 HTML 을 고쳐도 docx 는
옛 문장을 유지했다 — 확장 개수, 배정 결론, NP 귀속이 실제로 갈렸다.
원고를 고칠 곳은 HTML 하나뿐이다.
"""
import copy, glob, json, re, shutil, pathlib
from collections import Counter
from html import unescape
from docx import Document
from docx.shared import Pt, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import pathlib
ROOT = pathlib.Path(__file__).resolve().parents[1]   # 저장소 최상위

SRC = "/Users/otaegyeong/Downloads/워드템플릿.docx"
DST = str(ROOT / "paper/협동탐색_최적성_인증.docx")
FIG = str(ROOT / "paper/fig")
COL_TW = 4620          # 단 폭 [twips]
FRAME_W = 9639         # 제목부 프레임 폭 (2단 가로지르기)

shutil.copyfile(SRC, DST)
doc = Document(DST)
body = doc.element.body
sectPr = body.find(qn("w:sectPr"))
for child in list(body):
    if child is not sectPr:
        body.remove(child)

def _set(el, tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    el.append(e)
    return e

def para(text="", *, font="바탕", size=9, bold=False, align=None,
         indent=False, frame=False, line=280, space_after=0, tabs=None):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    if frame:
        f = OxmlElement("w:framePr")
        for k, v in (("w", FRAME_W), ("h", 301), ("wrap", "around"),
                     ("hAnchor", "page"), ("x", 1135), ("yAlign", "top")):
            f.set(qn(f"w:{k}"), str(v))
        pPr.append(f)
    _set(pPr, "w:spacing", line=line, lineRule="atLeast", after=space_after)
    if indent:
        _set(pPr, "w:ind", firstLineChars=100, firstLine=180)
    if align is not None:
        p.alignment = align
    if tabs:
        t = OxmlElement("w:tabs")
        for pos, kind in tabs:
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), kind); tab.set(qn("w:pos"), str(pos))
            t.append(tab)
        pPr.append(t)
    if text:
        r = p.add_run(text)
        r.font.name = font
        r.font.size = Pt(size)
        r.bold = bold
        rPr = r._element.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rPr.insert(0, rf)
        for a in ("ascii", "eastAsia", "hAnsi"):
            rf.set(qn(f"w:{a}"), font)
    return p

def eq(lines, number):
    """중앙 수식 + 우측 번호."""
    for i, line in enumerate(lines):
        p = para(align=WD_ALIGN_PARAGRAPH.CENTER, line=260,
                 tabs=[(COL_TW, "right")] if i == len(lines) - 1 else None)
        r = p.add_run(line); r.font.name = "바탕"; r.font.size = Pt(9); r.italic = True
        if i == len(lines) - 1:
            r2 = p.add_run(f"\t({number})")
            r2.font.name = "바탕"; r2.font.size = Pt(9)

def head(text):
    para(text, font="돋움", size=9, bold=True, line=280)

def figure(png, caption):
    # 그림은 앞뒤로 한 줄씩 띄운다. 본문에 바로 붙으면 2단 조판에서
    # 그림과 문단의 경계가 사라져 읽는 흐름이 끊긴다.
    para("", line=120)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(png, width=Inches(3.05))
    para(caption, font="돋움", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, line=260)
    para("", line=120)


# ================================================= HTML 원고 -> 블록 열
SOURCE = ROOT / "paper/협동탐색_최적성_인증_초고.html"
html = SOURCE.read_text(encoding="utf-8")


def _text(fragment: str) -> str:
    """태그를 걷어내고 실체참조를 푼다."""
    t = re.sub(r"<br\s*/?>", " ", fragment)
    t = re.sub(r"<[^>]+>", "", t)
    t = unescape(t)
    return re.sub(r"\s+", " ", t).strip()


# 표제부의 영문 제목은 HTML 에서 읽는다. 여기에 문자열로 적어 두었더니
# 제목을 바꾼 뒤에도 docx 만 옛 제목을 달고 나왔다.
TITLE_EN = _text(re.search(r'<p class="entitle">(.*?)</p>', html, re.S).group(1))


def blocks():
    """본문을 (종류, 내용) 열로 바꾼다. 순서는 HTML 그대로다."""
    body = html[html.find("<h1") :]
    # 목차(nav)와 표제부 부속 문단은 본문이 아니다. 걷어내지 않으면
    # 24 항목짜리 목차가 번호 목록으로, 영문 제목이 본문 첫 문단으로 딸려 나온다.
    body = re.sub(r"<nav\b.*?</nav>", "", body, flags=re.S)
    body = re.sub(r'<p class="(?:entitle|byline|kicker)">.*?</p>', "", body, flags=re.S)
    pattern = re.compile(
        r'<h([123])[^>]*>(?P<h>.*?)</h\1>'
        r'|<div class="eq">(?P<eq>.*?)</div>'
        r'|<div class="note">(?P<note>.*?)</div>'
        r'|<ol class="refs">(?P<refs>.*?)</ol>'
        r'|<ol[^>]*>(?P<ol>.*?)</ol>'
        r'|<figure>(?P<fig>.*?)</figure>'
        r'|<table>(?P<tab>.*?)</table>'
        r'|<p[^>]*>(?P<p>.*?)</p>',
        re.S,
    )
    for m in pattern.finditer(body):
        if m.group("h") is not None:
            yield ("h" + m.group(1), _text(m.group("h")))
        elif m.group("eq") is not None:
            raw = m.group("eq")
            lines = [l for l in _text(re.search(r"<pre>(.*?)</pre>", raw, re.S).group(1)).split("\n") if l]
            if not lines:
                lines = [_text(re.search(r"<pre>(.*?)</pre>", raw, re.S).group(1))]
            num = _text(re.search(r'<span class="tag">(.*?)</span>', raw, re.S).group(1)).strip("()")
            yield ("eq", (lines, num))
        elif m.group("note") is not None:
            yield ("note", _text(m.group("note")))
        elif m.group("refs") is not None:
            yield ("refs", [_text(x) for x in re.findall(r"<li>(.*?)</li>", m.group("refs"), re.S)])
        elif m.group("ol") is not None:
            yield ("ol", [_text(x) for x in re.findall(r"<li>(.*?)</li>", m.group("ol"), re.S)])
        elif m.group("fig") is not None:
            cap = re.search(r"<figcaption[^>]*>(.*?)</figcaption>", m.group("fig"), re.S)
            yield ("fig", _text(cap.group(1)) if cap else "")
        elif m.group("tab") is not None:
            raw = m.group("tab")
            cap = re.search(r"<caption>(.*?)</caption>", raw, re.S)
            head_cells = [_text(c) for c in re.findall(r"<th[^>]*>(.*?)</th>", raw, re.S)]
            body_rows = []
            for tr in re.findall(r"<tr>(.*?)</tr>", raw, re.S):
                cells = re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)
                if cells:
                    body_rows.append([_text(c) for c in cells])
            yield ("table", (_text(cap.group(1)) if cap else "", head_cells, body_rows))
        else:
            t = _text(m.group("p"))
            if t:
                yield ("p", t)


def figure_png(caption: str):
    """'그림 3. …' 또는 'Fig. 3. …' -> paper/fig/fig3_*.png . 없으면 건너뛴다."""
    m = re.match(r"(?:그림|Fig\.?)\s*(\d+)", caption)
    if not m:
        return None
    hits = sorted(pathlib.Path(FIG).glob(f"fig{m.group(1)}_*.png"))
    return str(hits[0]) if hits else None


def add_table(caption, head_cells, body_rows):
    para("", line=120)
    para(caption, font="돋움", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line=260)
    t = doc.add_table(rows=1, cols=len(head_cells))
    # 템플릿에 표 스타일이 없을 수 있다. 있으면 쓰고, 없으면 테두리를 직접 넣는다.
    try:
        t.style = "Table Grid"
    except KeyError:
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "808080")
            borders.append(e)
        tblPr.append(borders)
    for c, name in zip(t.rows[0].cells, head_cells):
        c.text = ""
        r = c.paragraphs[0].add_run(name)
        r.font.name = "돋움"; r.font.size = Pt(8); r.bold = True
    for row in body_rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = ""
            r = c.paragraphs[0].add_run(v)
            r.font.name = "바탕"; r.font.size = Pt(8)
    para("", line=120)


# ================================================= 제목부 (프레임)
title_ko = next((v for k, v in blocks() if k == "h1"), "제목")
para(title_ko, font="돋움", size=13, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, frame=True, line=260)
para(TITLE_EN, font="돋움", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, frame=True, line=260)
para("", frame=True, line=240)
para("홍길동*·성명2**", font="돋움", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, frame=True, line=240)
para("* 소속1   ** 소속2", font="돋움", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, frame=True, line=240)
para("(발표자 이메일 주소)", font="돋움", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, frame=True, line=240)
para("", frame=True, line=240)

# ================================================= 본문
skipped_figs = []
for kind, payload in blocks():
    if kind == "h1":
        continue
    if kind in ("h2", "h3"):
        head(payload)
    elif kind == "p":
        para(payload, indent=True)
    elif kind == "note":
        para(payload, size=8, indent=True)
    elif kind == "ol":
        for i, item in enumerate(payload, 1):
            para(f"{i}) {item}", indent=True)
    elif kind == "refs":
        for i, item in enumerate(payload, 1):
            para(f"[{i}] {item}", size=8)
    elif kind == "eq":
        eq(*payload)
    elif kind == "fig":
        png = figure_png(payload)
        if png:
            figure(png, payload)
        else:
            skipped_figs.append(payload[:48])
    elif kind == "table":
        add_table(*payload)

doc.save(DST)
counts = Counter(k for k, _ in blocks())
print(f"저장: {DST}")
print("  블록:", ", ".join(f"{k}={v}" for k, v in sorted(counts.items())))
if skipped_figs:
    print("  그림 파일 없어 건너뜀:", "; ".join(skipped_figs))
